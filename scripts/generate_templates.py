import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Directory setup
TEMPLATES_DIR = "data/templates"
os.makedirs(TEMPLATES_DIR, exist_ok=True)

# 1. Dimension Data
dimensions = {
    "dim_entity": {
        "columns": ["Entity", "Entity_Group", "Sort_Order"],
        "data": [
            ["CRDF", "Individual", 1],
            ["CAF", "Individual", 2],
            ["Combined", "Group", 3]
        ]
    },
    "dim_centre": {
        "columns": ["Centre", "Centre_Group", "Active_Flag", "Sort_Order"],
        "data": [
            ["CUPP", "Program", "Y", 1],
            ["CARBSE", "Program", "Y", 2],
            ["CWAS", "Program", "Y", 3],
            ["CEI", "Program", "Y", 4],
            ["CHC", "Program", "Y", 5],
            ["CAG", "Program", "Y", 6],
            ["CAU", "Program", "Y", 7],
            ["CoEUT", "Program", "Y", 8],
            ["HO", "Support", "Y", 9],
            ["Legacy", "Historical", "N", 10],
            ["Other", "Miscellaneous", "Y", 11]
        ]
    },
    "dim_segment": {
        "columns": ["Segment", "Segment_Group", "Sort_Order"],
        "data": [
            ["Core Advisory", "Operations", 1],
            ["PMU", "Operations", 2],
            ["Testing", "Operations", 3],
            ["Conference", "Operations", 4],
            ["Grant", "Funding", 5],
            ["Centre Aggregate", "Reporting", 6],
            ["Entity Aggregate", "Reporting", 7]
        ]
    },
    "dim_revenue_type": {
        "columns": ["Revenue_Type", "Revenue_Group", "Sort_Order"],
        "data": [
            ["Earned Fees", "Operating", 1],
            ["Grant", "Operating", 2],
            ["Fund Transfer", "Non-Operating", 3],
            ["Total Turnover", "Total", 4]
        ]
    },
    "dim_record_type": {
        "columns": ["Record_Type", "Description", "Sort_Order"],
        "data": [
            ["Segment_Row", "Data at segment level", 1],
            ["Centre_Row", "Data at centre level", 2],
            ["Entity_Row", "Data at entity level", 3]
        ]
    }
}

def generate_dim_month():
    start_date = "2024-04-01"
    end_date = "2028-03-31"
    dates = pd.date_range(start=start_date, end=end_date, freq='ME')

    month_data = []
    for i, dt in enumerate(dates):
        month_label = dt.strftime('%b-%y')
        # Indian FY: Apr to Mar.
        # If month >= 4, FY is current_year-next_year
        # If month < 4, FY is prev_year-current_year
        if dt.month >= 4:
            fy = f"FY{dt.strftime('%y')}-{str(dt.year + 1)[2:]}"
            month_no = dt.month - 3
            quarter = (dt.month - 4) // 3 + 1
        else:
            fy = f"FY{str(dt.year - 1)[2:]}-{dt.strftime('%y')}"
            month_no = dt.month + 9
            quarter = 4

        month_data.append([
            dt.strftime('%Y-%m-%d'),
            month_label,
            fy,
            f"Q{quarter}",
            month_no,
            i + 1
        ])

    df = pd.DataFrame(month_data, columns=["Month", "Month_Label", "FY", "Quarter", "Month_No", "Sort_Order"])
    df.to_csv(os.path.join(TEMPLATES_DIR, "dim_month.csv"), index=False)
    return df

# 2. Fact Table Columns
fact_columns = [
    "Month", "Month_Label", "FY", "Period_No", "Entity", "Centre", "Segment",
    "Revenue_Type", "Record_Type", "Revenue", "Direct_Cost", "Gross_Margin",
    "Overhead_CEPT", "Overhead_Internal", "Total_Overhead", "Provision_Debtors",
    "Provision_Other", "Operating_Surplus", "Tax", "Net_Surplus", "AR_Closing",
    "AR_0_30", "AR_31_60", "AR_61_90", "AR_91_180", "AR_181_365", "AR_365_Plus",
    "WIP_Closing", "Collections", "Notes", "Source_Reference"
]

def main():
    # Generate Dimension CSVs
    for name, info in dimensions.items():
        df = pd.DataFrame(info["data"], columns=info["columns"])
        df.to_csv(os.path.join(TEMPLATES_DIR, f"{name}.csv"), index=False)

    generate_dim_month()

    # Generate Fact Template CSV
    df_fact = pd.DataFrame(columns=fact_columns)
    df_fact.to_csv(os.path.join(TEMPLATES_DIR, "fact_financials_template.csv"), index=False)

    # Generate Fact Template XLSX
    wb = Workbook()

    # Sheet 1: Fact_Financials
    ws_fact = wb.active
    ws_fact.title = "Fact_Financials"
    ws_fact.append(fact_columns)
    ws_fact.freeze_panes = "A2"

    # Sheet 2: Lists
    ws_lists = wb.create_sheet("Lists")
    list_cols = ["Entity", "Centre", "Segment", "Revenue_Type", "Record_Type"]
    ws_lists.append(list_cols)

    # Populate Lists sheet
    max_rows = 0
    for i, col in enumerate(list_cols):
        dim_key = f"dim_{col.lower()}"
        values = [row[0] for row in dimensions[dim_key]["data"]]
        for j, val in enumerate(values):
            ws_lists.cell(row=j+2, column=i+1, value=val)
        max_rows = max(max_rows, len(values))

    # Add Data Validation to Fact_Financials
    # Dropdowns for Entity (E), Centre (F), Segment (G), Revenue_Type (H), Record_Type (I)
    validation_map = {
        "E": ("Lists!$A$2:$A$100", "Entity"),
        "F": ("Lists!$B$2:$B$100", "Centre"),
        "G": ("Lists!$C$2:$C$100", "Segment"),
        "H": ("Lists!$D$2:$D$100", "Revenue_Type"),
        "I": ("Lists!$E$2:$E$100", "Record_Type")
    }

    for col_letter, (formula, title) in validation_map.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = f'Please select a value from the {title} list.'
        dv.errorTitle = 'Invalid Selection'
        ws_fact.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}1000") # Apply to first 1000 rows

    # Sheet 3: Instructions
    ws_instr = wb.create_sheet("Instructions")
    instructions = [
        ["Guidance"],
        ["- one row = one record"],
        ["- use standard values only"],
        ["- do not merge cells"],
        ["- do not insert formulas into fact rows"],
        ["- keep all numeric cells numeric"],
        ["- leave blanks rather than text like 'NA'"]
    ]
    for row in instructions:
        ws_instr.append(row)

    wb.save(os.path.join(TEMPLATES_DIR, "fact_financials_template.xlsx"))
    print("Templates generated successfully.")

if __name__ == "__main__":
    main()
